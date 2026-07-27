"""站级案例钉扎 + Word 报告锁测(功能② 批次 4:证据链出口)。

语义:导出必经钉扎(项目报告 auto-pin 同律)——station_cases 冻结「站型文档 + 量化
结果 + doc_hash + 引擎版本」,站型后续修改不追溯已钉案例;docx 从钉扎案例生成,
KPI 标签绑定(≥5 舱 ↔ ge5,整排 ↔ ge_row)在文档里同样成立;口径申报与免责逐字入文;
站级签发通道未建 ⇒ 恒工作稿标记。
"""
import io
import json

import pytest
from fastapi.testclient import TestClient

from backend import store
from backend.app import app
from tests.backend.test_station_model import _mk_station

client = TestClient(app)


def _create_station() -> int:
    r = client.post("/api/station", json=_mk_station())
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_pin_freezes_doc_and_results_with_audit():
    sid = _create_station()
    r = client.post(f"/api/station/{sid}/pin")
    assert r.status_code == 200, r.text
    case = r.json()
    assert case["station_id"] == sid
    assert len(case["doc_hash"]) == 64
    assert case["results"]["metrics"]["ge5"] == 0.0                 # 3 舱排:literal ≥5 不可达
    assert case["engine_version"]
    # 列表可见
    lst = client.get(f"/api/station/{sid}/cases").json()
    assert any(c["id"] == case["id"] for c in lst)
    # 不可变性:站型改名/改参后,已钉案例的 doc 与 hash 不动
    doc = client.get(f"/api/station/{sid}").json()
    doc["name"] = "改名后"
    doc["edge_defaults"]["intra_p"] = 0.5
    assert client.put(f"/api/station/{sid}", json=doc).status_code == 200
    again = client.get(f"/api/station/case/{case['id']}").json()
    assert again["doc_hash"] == case["doc_hash"]
    assert again["doc"]["edge_defaults"]["intra_p"] == pytest.approx(0.10)
    # 审计
    with store._conn() as cx:
        acts = [x["action"] for x in cx.execute(
            "SELECT action FROM audit_log WHERE action='station.case.pin'").fetchall()]
    assert acts, "钉扎必须留审计"


def test_report_docx_contents_and_label_binding():
    sid = _create_station()
    cid = client.post(f"/api/station/{sid}/pin").json()["id"]
    r = client.get(f"/api/station/case/{cid}/report.docx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    from docx import Document
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    allt = text + "\n" + cells
    # 标签绑定:两个指标名与各自数值同现,且值不同(3 舱站 ge5=0,ge_row>0)
    assert "≥5 舱(排内连片)" in allt and "整排蔓延(k=排长)" in allt
    q = client.get(f"/api/station/case/{cid}").json()["results"]["metrics"]
    assert f"{q['ge_row']:.2e}" in allt
    assert f"{q['ge5']:.2e}" in allt and q["ge5"] != q["ge_row"]
    # 证据链要素
    assert "doc_hash" in allt or "站型文档哈希" in allt
    case = client.get(f"/api/station/case/{cid}").json()
    assert case["doc_hash"][:8] in allt                      # 哈希短码入文
    # 产品版本哈希(舱=版本快照实例)
    vh = case["results"]["products"][0]["case"]["version_hash"][:8]
    assert vh in allt
    # 口径申报逐字 + 工作稿 + 免责
    assert "点值未计入跨舱 CCF" in allt
    assert "工作稿" in allt
    assert "正式工程分析必须" in allt or "参考量级" in allt
    # 站级签发通道未建的诚实声明
    assert "站级签发" in allt


def test_pin_missing_station_and_dangling_member_are_loud():
    assert client.post("/api/station/99999/pin").status_code == 404
    sid = _create_station()
    body = client.get(f"/api/station/{sid}").json()
    with store._conn() as cx:
        cx.execute("DELETE FROM analysis_cases WHERE id=?", (body["products"][0]["case_id"],))
    r = client.post(f"/api/station/{sid}/pin")
    assert r.status_code == 404 and str(body["products"][0]["case_id"]) in r.text
    assert client.get("/api/station/case/99999").status_code == 404
    assert client.get("/api/station/case/99999/report.docx").status_code == 404
