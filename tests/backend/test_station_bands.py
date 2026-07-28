"""站级不确定带锁测(规格 2026-07-27-station-bands,ST-6)。

语义:同分位代入 = 全联动包络(公式单调 ⇒ 输出分位精确/带序不倒挂);尾和分位 =
Σ 箱 band[q],点值>0 无带的箱以点值参与并响亮记账;全无带 ⇒ bands null + note;
倒挂带 422;点值通道逐位不变。
"""
import pytest
from fastapi.testclient import TestClient

from backend.app import app
from backend import store
from backend.canonical import to_canonical
from tests.backend.test_compare_assemble import _payload, PA, PB

client = TestClient(app)

# 规格 §6 夹具带
BANDS_A = {"C3": {"p05": 5e-5, "p50": 1.1e-4, "p95": 3e-4},
           "C4": {"p05": 4e-6, "p50": 1e-5, "p95": 4e-5}}
BANDS_B = {"C4": {"p05": 8e-6, "p50": 2e-5, "p95": 6e-5}}
CONS_A = {"C1": 7e-4, "C2": 5e-5, "C3": 1e-4, "C4": 1e-5, "C5": 0.0}
CONS_B = {"C1": 9e-4, "C2": 3e-5, "C3": 2e-4, "C4": 2e-5, "C5": 0.0}


def _pin_with(puuid: str, name: str, cons: dict, bands: dict | None) -> int:
    store.ensure_project(puuid, name)
    vid = store.create_version(puuid, *to_canonical(_payload(0.10), puuid))["id"]
    results: dict = {"consequences": cons}
    if bands is not None:
        results["bands"] = bands
    return store.create_case(puuid, vid, "report", results=results)


def _station_body(ca: int, cb: int) -> dict:
    return {
        "name": "带锚站",
        "products": [{"key": "A", "label": "供应商A", "project_uuid": PA, "case_id": ca},
                     {"key": "B", "label": "供应商B", "project_uuid": PB, "case_id": cb}],
        "rows": [{"id": "Z1", "label": "A 排", "cabins": [
            {"id": "A-01", "product": "A"}, {"id": "A-02", "product": "A"},
            {"id": "A-03", "product": "B"}], "walls": []}],
        "edge_defaults": {"intra_p": 0.10, "wall_p": 0.03},
        "edge_overrides": [],
        "station_barriers": [
            {"key": "sup", "label": "站级消防", "pfd": 0.5, "grade": "D", "scope": "全站", "basis": "测试"},
            {"key": "er", "label": "应急响应", "pfd": 0.5, "grade": "D", "scope": "全站", "basis": "测试"}],
        "ccf_groups": [],
        "offsite": {"p_fence": 0.2, "basis": "测试"},
    }


def _quantify(body: dict) -> dict:
    r = client.post("/api/station", json=body)
    assert r.status_code == 200, r.text
    q = client.get(f"/api/station/{r.json()['id']}/quantify")
    assert q.status_code == 200, q.text
    return q.json()


def test_B1_band_composition_hand_values_ordering_and_point_untouched():
    ca = _pin_with(PA, "供应商A", CONS_A, BANDS_A)
    cb = _pin_with(PB, "供应商B", CONS_B, BANDS_B)
    d = _quantify(_station_body(ca, cb))
    b = d["bands"]
    # 规格 §6 表(同分位合成手算)
    assert b["any"]["p05"] == pytest.approx(3.159746e-4, rel=1e-6)
    assert b["any"]["p95"] == pytest.approx(9.397076e-4, rel=1e-6)
    assert b["ge2"]["p05"] == pytest.approx(1.823000e-5, rel=1e-9)
    assert b["ge2"]["p50"] == pytest.approx(2.840000e-5, rel=1e-9)
    assert b["ge2"]["p95"] == pytest.approx(6.230000e-5, rel=1e-9)
    assert b["ge_row"]["p05"] == pytest.approx(7.9e-7, rel=1e-9)
    assert b["ge_row"]["p95"] == pytest.approx(2.35e-6, rel=1e-9)
    assert b["offsite"]["p95"] == pytest.approx(2.799872e-5, rel=1e-6)
    # 3 舱排 ge5 恒 0:带亦 0(literal ≥5 排内不可达)
    assert b["ge5"]["p95"] == 0.0
    # 带序不倒挂(全指标)
    for k in ("any", "ge2", "ge5", "ge_row", "offsite"):
        assert b[k]["p05"] <= b[k]["p50"] <= b[k]["p95"], k
    # 点值通道逐位不变(v1 闭式,与 test_station_model 同式)
    f = [1.1e-4, 1.1e-4, 2.2e-4]
    esc = [0.10, 0.19, 0.10]
    assert d["metrics"]["ge2"] == pytest.approx(sum(fi * e for fi, e in zip(f, esc)) * 0.5, rel=1e-9)
    assert d["metrics"]["any"] == pytest.approx(1 - (1 - f[0]) * (1 - f[1]) * (1 - f[2]), rel=1e-9)
    # band_status:A 参与箱(C3,C4)全有带 = full;B 的 C3 点值>0 无带 = partial
    st = {p["key"]: p["band_status"] for p in d["products"]}
    assert st == {"A": "full", "B": "partial"}
    assert any("B" in n and "C3" in n for n in d["band_notes"]), d["band_notes"]
    # 产品尾和带外露(检视器用)
    pa = next(p for p in d["products"] if p["key"] == "A")
    assert pa["f_band"]["p95"] == pytest.approx(3.4e-4, rel=1e-9)


def test_B2_all_products_bandless_yields_null_plus_note():
    ca = _pin_with(PA, "供应商A", CONS_A, None)
    cb = _pin_with(PB, "供应商B", CONS_B, None)
    d = _quantify(_station_body(ca, cb))
    assert d["bands"] is None
    assert any("无不确定带" in n or "无带" in n for n in d["band_notes"]), d["band_notes"]
    assert all(p["band_status"] == "none" for p in d["products"])
    assert all(p["f_band"] is None for p in d["products"])


def test_B3_inverted_band_rejected_loudly():
    bad = {"C4": {"p05": 6e-5, "p50": 2e-5, "p95": 8e-6}}      # 倒挂
    ca = _pin_with(PA, "供应商A", CONS_A, bad)
    cb = _pin_with(PB, "供应商B", CONS_B, None)
    r = client.post("/api/station", json=_station_body(ca, cb))
    assert r.status_code == 200, r.text
    q = client.get(f"/api/station/{r.json()['id']}/quantify")
    assert q.status_code == 422
    assert "A" in q.text and "C4" in q.text                     # 点名产品与箱


def test_B4_report_docx_band_columns_when_present():
    import io
    from docx import Document
    ca = _pin_with(PA, "供应商A", CONS_A, BANDS_A)
    cb = _pin_with(PB, "供应商B", CONS_B, BANDS_B)
    r = client.post("/api/station", json=_station_body(ca, cb))
    sid = r.json()["id"]
    cid = client.post(f"/api/station/{sid}/pin").json()["id"]
    dx = client.get(f"/api/station/case/{cid}/report.docx")
    assert dx.status_code == 200
    doc = Document(io.BytesIO(dx.content))
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "P05" in cells and "P95" in cells
    assert f"{6.230000e-5:.2e}" in cells                        # ge2 p95 入表
    assert "C3" in text and "点值" in text                       # band_notes 段(B 部分带)
