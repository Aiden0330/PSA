"""站级验收准则判定锁测(规格 2026-07-27-station-acceptance,ST-7)。

判定是监管口径声明:判据口径显式必填、限值必须带依据、无带不可判 ⇒ 响亮
indeterminate(绝不静默降级点值判)。pass ⇔ value ≤ limit;margin = value/limit。
"""
import pytest
from fastapi.testclient import TestClient

from backend.app import app
from tests.backend.test_station_bands import (
    BANDS_A, BANDS_B, CONS_A, CONS_B, _pin_with, _station_body)
from tests.backend.test_compare_assemble import PA, PB

client = TestClient(app)


def _mk(acceptance, bands_a=None, bands_b=None) -> dict:
    ca = _pin_with(PA, "供应商A", CONS_A, bands_a)
    cb = _pin_with(PB, "供应商B", CONS_B, bands_b)
    body = _station_body(ca, cb)
    body["acceptance"] = acceptance
    return body


def _quantify(body: dict, expect: int = 200):
    r = client.post("/api/station", json=body)
    assert r.status_code == 200, r.text
    q = client.get(f"/api/station/{r.json()['id']}/quantify")
    assert q.status_code == expect, q.text
    return q


def test_A1_point_verdicts_pass_fail_margin_order():
    d = _quantify(_mk([
        {"metric": "ge_row", "limit": 1e-5, "eval_basis": "point", "basis": "业主目标(示例)"},
        {"metric": "ge2", "limit": 1e-5, "eval_basis": "point", "basis": "内部准则(示例)", "label": "多舱目标"},
    ])).json()
    ar = d["acceptance_results"]
    assert [a["metric"] for a in ar] == ["ge_row", "ge2"]        # 顺序 = 文档序
    assert ar[0]["verdict"] == "pass"
    assert ar[0]["value"] == pytest.approx(1.1e-6, rel=1e-9)
    assert ar[0]["margin"] == pytest.approx(0.11, rel=1e-9)
    assert ar[1]["verdict"] == "fail"
    assert ar[1]["margin"] == pytest.approx(2.695, rel=1e-9)
    assert ar[1]["label"] == "多舱目标" and ar[1]["basis"]


def test_A2_p95_verdicts_use_band_value():
    body = _mk([
        {"metric": "ge2", "limit": 1e-4, "eval_basis": "p95", "basis": "保守判据(示例)"},
        {"metric": "ge2", "limit": 5e-5, "eval_basis": "p95", "basis": "保守判据(示例)"},
    ], bands_a=BANDS_A, bands_b=BANDS_B)
    ar = _quantify(body).json()["acceptance_results"]
    assert ar[0]["verdict"] == "pass" and ar[0]["value"] == pytest.approx(6.23e-5, rel=1e-9)
    assert ar[1]["verdict"] == "fail"


def test_A3_p95_without_bands_is_loud_indeterminate():
    body = _mk([
        {"metric": "ge2", "limit": 1e-4, "eval_basis": "p95", "basis": "保守判据(示例)"},
        {"metric": "ge2", "limit": 1e-4, "eval_basis": "point", "basis": "点值判据(示例)"},
    ])                                                            # 无带
    ar = _quantify(body).json()["acceptance_results"]
    assert ar[0]["verdict"] == "indeterminate"
    assert ar[0]["value"] is None and ar[0]["margin"] is None
    assert "无带" in (ar[0].get("reason") or "")
    assert ar[1]["verdict"] == "pass"                             # 同站 point 准则正常判


def test_A4_validation_loud_422s():
    cases = [
        ({"metric": "ge9", "limit": 1e-5, "eval_basis": "point", "basis": "x"}, "ge9"),
        ({"metric": "ge2", "limit": 0, "eval_basis": "point", "basis": "x"}, "limit"),
        ({"metric": "ge2", "limit": 1e-5, "eval_basis": "point", "basis": "  "}, "basis"),
        ({"metric": "ge2", "limit": 1e-5, "basis": "x"}, "eval_basis"),
    ]
    for acc, needle in cases:
        r = client.post("/api/station", json=_mk([acc]))
        assert r.status_code == 422, f"{acc} 应 422"
        assert needle in r.text


def test_A5_docx_acceptance_chapter():
    import io
    from docx import Document
    body = _mk([
        {"metric": "ge5", "limit": 1e-6, "eval_basis": "point", "basis": "NFPA 855 场址目标(示例)"},
        {"metric": "ge2", "limit": 1e-5, "eval_basis": "point", "basis": "内部准则(示例)"},
    ])
    r = client.post("/api/station", json=body)
    sid = r.json()["id"]
    cid = client.post(f"/api/station/{sid}/pin").json()["id"]
    dx = client.get(f"/api/station/case/{cid}/report.docx")
    assert dx.status_code == 200
    doc = Document(io.BytesIO(dx.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    allt = text + "\n" + cells
    assert "验收准则判定" in allt
    assert "PASS" in allt and "FAIL" in allt                      # ge5 过(0≤1e-6)、ge2 不过
    assert "NFPA 855 场址目标(示例)" in allt                     # basis 落文
    assert "≥5 舱(排内连片)" in cells                          # 判定表用 canonical 标签
    assert "存在未通过准则" in text                               # FAIL 章首声明
